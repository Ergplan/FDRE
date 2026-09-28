from __future__ import annotations

import json
import sys
from pathlib import Path

import docx


def classify_paragraph(text: str, style: str) -> dict:
    upper = text.strip().upper()
    if not text.strip():
        return {}
    if upper.startswith("TABLE "):
        return {"type": "caption", "text": text.strip(), "style": style}
    if upper.startswith("NOTE"):
        return {"type": "note_label", "text": text.strip(), "style": style}
    if style.lower().startswith("heading"):
        level = 2
        for token in style.split():
            if token.isdigit():
                level = int(token)
                break
        return {"type": "heading", "level": level, "text": text.strip(), "style": style}
    if upper[:2].isdigit() or upper[:2] in {"1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9."}:
        if len(text.strip()) < 120 and any(ch.isalpha() for ch in text):
            return {"type": "heading", "level": 2, "text": text.strip(), "style": style}
    return {"type": "paragraph", "text": text.strip(), "style": style}


def table_block(table, caption: str | None) -> dict:
    rows = []
    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])
    header = rows[0] if rows else []
    return {
        "type": "table",
        "caption": caption or "",
        "rows": rows,
        "header": header,
        "body": rows[1:] if len(rows) > 1 else [],
        "column_count": max((len(row) for row in rows), default=0),
    }


def ingest_docx(path: Path) -> dict:
    doc = docx.Document(str(path))
    blocks = []
    pending_caption = None

    body = doc.element.body
    para_map = {p._p: p for p in doc.paragraphs}
    table_map = {t._tbl: t for t in doc.tables}

    for child in body.iterchildren():
        if child in para_map:
            para = para_map[child]
            text = para.text.strip()
            block = classify_paragraph(text, para.style.name if para.style else "Normal")
            if not block:
                continue
            if block["type"] == "caption":
                pending_caption = block["text"]
                blocks.append(block)
            else:
                blocks.append(block)
        elif child in table_map:
            block = table_block(table_map[child], pending_caption)
            blocks.append(block)
            pending_caption = None

    paragraphs = [b for b in blocks if b.get("type") in {"paragraph", "heading"}]
    title = paragraphs[0]["text"] if paragraphs else path.stem
    subtitle = paragraphs[1]["text"] if len(paragraphs) > 1 else ""
    headings = [b["text"] for b in blocks if b.get("type") == "heading"]
    return {
        "title": title,
        "subtitle": subtitle,
        "source_name": path.name,
        "block_count": len(blocks),
        "image_count": len(doc.inline_shapes),
        "table_count": sum(1 for b in blocks if b.get("type") == "table"),
        "headings": headings,
        "blocks": blocks,
    }


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: ingest_joulewise_docx.py input.docx output.json")
    source = Path(sys.argv[1])
    target = Path(sys.argv[2])
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(ingest_docx(source), indent=2, ensure_ascii=False), encoding="utf-8")


if __name__ == "__main__":
    main()
