from __future__ import annotations

import io
import math
import os
import re
import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
from fdre_tender_review import attach_review, document_role


try:
    from pypdf import PdfReader  # type: ignore
except Exception:  # pragma: no cover - optional dependency
    PdfReader = None


try:
    import fitz  # type: ignore
except Exception:  # pragma: no cover - optional dependency
    fitz = None


try:
    import docx  # type: ignore
except Exception:  # pragma: no cover - optional dependency
    docx = None


AMENDMENT_KEYWORDS = ("amendment", "corrigendum", "addendum", "clarification")
CHANGE_REF_RE = re.compile(
    r"(?i)\b(?:clause|section|annexure|appendix|format)\s+([a-z0-9ivx]+(?:\.[a-z0-9]+)*)"
)
DATE_RE = re.compile(
    r"(?i)\b(?:\d{1,2}[./-]\d{1,2}[./-]\d{2,4}|"
    r"\d{1,2}\s+(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+\d{2,4})\b"
)
CLAUSE_RE = re.compile(
    r"(?im)(?:^|[\n\f])\s*(?:#{1,6}\s+)?"
    r"((?:\d{1,2}\.)+\d{1,3}|section\s+\d+[a-z]?|annexure\s*[-–]?\s*[a-z0-9ivx]+|"
    r"appendix\s*[-–]?\s*[a-z0-9ivx]+|format\s*\d+(?:\.\d+)*)"
    r"\b[\s:.\-–]*(.{0,160})"
)

QUERY_EXPANSIONS = {
    "tender summary": "procurement capacity tender size quantum supply selection rfs bidder project location ISTS ppa",
    "key dates": "schedule last date bid submission pre bid meeting reverse auction commissioning scod ppa execution",
    "financial conditions": "emd earnest money pbg performance bank guarantee processing fee success charge payment security psm tariff penalty liquidated damages net worth turnover line of credit",
    "technical conditions": "peak availability cuf firm dispatchable storage bess ess generation schedule contracted capacity",
    "amendments": "amendment corrigendum addendum clarification clause modified deleted added",
}


@dataclass
class TenderChunk:
    content: str
    meta: dict[str, Any]
    score: float = 0.0


def extract_text_from_upload(name: str, payload: bytes) -> tuple[str, list[dict[str, Any]]]:
    suffix = Path(name).suffix.lower()
    pages: list[dict[str, Any]] = []
    if suffix == ".pdf":
        text = _extract_pdf_text(payload, name, pages)
    elif suffix == ".docx":
        text = _extract_docx_text(payload, name, pages)
    else:
        text = payload.decode("utf-8", errors="replace")
        pages.append({"page": 1, "text": text, "source": name})
    return text, pages


def _extract_docling(name: str, payload: bytes, do_ocr: bool = True) -> tuple[str, list[dict[str, Any]]]:
    from docling.datamodel.base_models import DocumentStream, InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption

    options = PdfPipelineOptions()
    options.do_ocr = do_ocr
    options.do_table_structure = True
    options.enable_remote_services = False
    options.document_timeout = 180
    # models baked into the engine image (docling-tools models download), so no runtime download
    artifacts = os.environ.get("DOCLING_ARTIFACTS_PATH")
    if artifacts and Path(artifacts).is_dir() and any(Path(artifacts).iterdir()):
        options.artifacts_path = artifacts
    converter = DocumentConverter(
        allowed_formats=[InputFormat.PDF, InputFormat.DOCX],
        format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)},
    )
    result = converter.convert(DocumentStream(name=Path(name).name, stream=io.BytesIO(payload)))
    if getattr(result.status, "value", result.status) != "success":
        raise RuntimeError("Docling did not complete the entire document.")
    document = result.document
    if Path(name).suffix.lower() == ".pdf":
        # Preserve physical PDF page boundaries for downstream source citations.
        pages = [
            {"page": page, "text": document.export_to_markdown(page_no=page), "source": name}
            for page in sorted(document.pages)
        ]
    else:
        pages = [{"page": 1, "text": document.export_to_markdown(), "source": name}]
    text = "\f".join(page["text"] for page in pages)
    if not text.strip():
        raise RuntimeError("Docling returned no readable text.")
    return text, pages


def extract_tender_document(name: str, payload: bytes, parser: str = "auto"):
    if parser not in {"auto", "standard", "docling"}:
        raise ValueError("Unknown tender parser.")
    warning = None
    if parser != "standard" and Path(name).suffix.lower() in {".pdf", ".docx"}:
        ocr_note = None
        try:
            try:
                text, pages = _extract_docling(name, payload)
            except ImportError:
                raise
            except Exception:
                # OCR models may be unavailable (no internet for the first download): layout and
                # table recognition still work for digital PDFs without OCR
                logging.getLogger(__name__).warning("Docling with OCR failed; retrying without OCR", exc_info=True)
                text, pages = _extract_docling(name, payload, do_ocr=False)
                ocr_note = "Docling ran without OCR (its OCR models could not be loaded): scanned pages have no text."
            if Path(name).suffix.lower() == ".pdf":
                native_pages = []
                try:
                    _extract_pdf_text(payload, name, native_pages)
                    for page, native in zip(pages, native_pages):
                        page["native_text"] = native["text"]
                except Exception:
                    pass
            table_note = "Review table contents against the original PDF; reconstructed cells may omit or repeat text." if Path(name).suffix.lower() == ".pdf" else None
            return text, pages, {"requested": parser, "engine": "Docling" if not ocr_note else "Docling (no OCR)", "warning": " ".join(x for x in (ocr_note, table_note) if x) or None}
        except ImportError:
            warning = "Docling is unavailable. Standard extraction was used. Install requirements-docling.txt to enable enhanced parsing."
        except Exception:
            logging.getLogger(__name__).exception("Docling extraction failed")
            warning = "Docling could not complete this document. Standard extraction was used; review tables and scanned pages."
    text, pages = extract_text_from_upload(name, payload)
    if not text.strip():
        raise ValueError("No readable text was extracted. For a scanned PDF, install Docling and retry with OCR enabled.")
    if any(not page["text"].strip() for page in pages):
        warning = (warning or "") + " Some pages contain no extracted text; review them for scanned content."
    return text, pages, {"requested": parser, "engine": "Standard", "warning": warning}


def _extract_pdf_text(payload: bytes, name: str, pages: list[dict[str, Any]]) -> str:
    if fitz is not None:
        document = fitz.open(stream=payload, filetype="pdf")
        all_pages = []
        for idx, page in enumerate(document, 1):
            page_text = page.get_text("text") or ""
            pages.append({"page": idx, "text": page_text, "source": name})
            all_pages.append(page_text)
        document.close()
        return "\f".join(all_pages)
    if PdfReader is None:
        raise RuntimeError("PDF text extraction requires PyMuPDF (fitz) or pypdf.")
    reader = PdfReader(io.BytesIO(payload))
    all_pages = []
    for idx, page in enumerate(reader.pages, 1):
        page_text = page.extract_text() or ""
        pages.append({"page": idx, "text": page_text, "source": name})
        all_pages.append(page_text)
    return "\f".join(all_pages)


def _extract_docx_text(payload: bytes, name: str, pages: list[dict[str, Any]]) -> str:
    if docx is None:
        return payload.decode("utf-8", errors="replace")
    document = docx.Document(io.BytesIO(payload))
    blocks: list[str] = []
    for para in document.paragraphs:
        if para.text.strip():
            blocks.append(para.text)
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                blocks.append(" | ".join(values))
    text = "\n".join(blocks)
    pages.append({"page": 1, "text": text, "source": name})
    return text


def normalize_space(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def clean_clause_id(value: str) -> str:
    return normalize_space(value).strip(" .:-–").upper()


def clean_clause_title(value: str) -> str:
    return normalize_space(value).strip(" .:-–")[:140]


def split_tender_chunks(name: str, text: str, max_words: int = 420, overlap: int = 70) -> list[TenderChunk]:
    chunks: list[TenderChunk] = []
    pages = text.split("\f") if "\f" in text else [text]
    for page_number, page_text in enumerate(pages, 1):
        page_chunks = _split_page(name, page_text, page_number, max_words, overlap)
        chunks.extend(page_chunks)
    for idx, chunk in enumerate(chunks):
        chunk.meta["split_id"] = idx
    return chunks


def _split_page(name: str, page_text: str, page_number: int, max_words: int, overlap: int) -> list[TenderChunk]:
    matches = list(CLAUSE_RE.finditer(page_text))
    if not matches:
        return _word_chunks(name, page_text, page_number, None, None, max_words, overlap)

    chunks: list[TenderChunk] = []
    if matches[0].start() > 0:
        chunks.extend(_word_chunks(name, page_text[: matches[0].start()], page_number, None, None, max_words, overlap))
    for idx, match in enumerate(matches):
        clause_id = clean_clause_id(match.group(1))
        clause_title = clean_clause_title(match.group(2))
        end = matches[idx + 1].start() if idx + 1 < len(matches) else len(page_text)
        chunks.extend(_word_chunks(name, page_text[match.start():end], page_number, clause_id, clause_title, max_words, overlap))
    return chunks


def _word_chunks(
    name: str,
    text: str,
    page_number: int,
    clause_id: str | None,
    clause_title: str | None,
    max_words: int,
    overlap: int,
) -> list[TenderChunk]:
    words = text.split()
    if not words:
        return []
    step = max(1, max_words - overlap)
    chunks: list[TenderChunk] = []
    for start in range(0, len(words), step):
        content = " ".join(words[start:start + max_words])
        if not content:
            continue
        meta = {
            "file_name": name,
            "page_number": page_number,
            "split_idx_start": start,
        }
        if clause_id:
            meta["clause_id"] = clause_id
        if clause_title:
            meta["clause_title"] = clause_title
        chunks.append(TenderChunk(content=content, meta=meta))
        if start + max_words >= len(words):
            break
    return chunks


def expand_query(query: str) -> str:
    lowered = query.lower()
    additions = []
    for key, terms in QUERY_EXPANSIONS.items():
        if any(token in lowered for token in key.split()):
            additions.append(terms)
    if "capacity" in lowered and "file" not in lowered:
        additions.append("procurement capacity contracted capacity bid capacity MW quantum")
    return normalize_space(" ".join([query, *additions]))


def retrieve_chunks(chunks: Iterable[TenderChunk], query: str, top_k: int = 6) -> list[TenderChunk]:
    expanded = expand_query(query)
    query_terms = _terms(expanded)
    if not query_terms:
        return []
    chunk_list = list(chunks)
    doc_freq: dict[str, int] = {}
    chunk_terms = []
    for chunk in chunk_list:
        terms = _terms(chunk.content)
        chunk_terms.append(terms)
        for term in set(terms):
            doc_freq[term] = doc_freq.get(term, 0) + 1
    total_docs = max(len(chunk_list), 1)
    scored: list[TenderChunk] = []
    for chunk, terms in zip(chunk_list, chunk_terms):
        counts: dict[str, int] = {}
        for term in terms:
            counts[term] = counts.get(term, 0) + 1
        length_norm = max(len(terms), 1)
        score = 0.0
        for term in query_terms:
            if term not in counts:
                continue
            idf = math.log(1 + (total_docs - doc_freq.get(term, 0) + 0.5) / (doc_freq.get(term, 0) + 0.5))
            score += idf * (counts[term] / math.sqrt(length_norm))
        if score > 0:
            scored.append(TenderChunk(chunk.content, dict(chunk.meta), score))
    scored.sort(key=lambda item: item.score, reverse=True)
    return scored[:top_k]


def _terms(text: str) -> list[str]:
    stop = {
        "the", "and", "for", "with", "shall", "this", "that", "from", "into", "under",
        "are", "will", "not", "may", "per", "has", "have", "was", "were", "been",
    }
    return [term for term in re.findall(r"[a-z0-9]{2,}", text.lower()) if term not in stop]


def parse_tender_document(name: str, payload: bytes, defaults: Any | None = None, parser: str = "auto") -> dict[str, Any]:
    text, _pages, extraction = extract_tender_document(name, payload, parser)
    chunks = split_tender_chunks(name, text)
    normalized = normalize_space(text)
    lower = normalized.lower()
    default_tender = getattr(defaults, "tender", None)

    quantum = _procurement_quantum(normalized)
    tender_size = quantum.get("procurement_mw") or _parse_tender_size(normalized, default_value=getattr(default_tender, "max_bidder_capacity_mw", 1200.0))
    ppa_years = _parse_int_context(normalized, ["ppa", "power purchase"], default=getattr(default_tender, "ppa_years", 25))
    parsed_peak_pct = _parse_percent_context(normalized, ["peak availability", "availability during peak"], default=None)
    peak_pct = parsed_peak_pct if parsed_peak_pct is not None else getattr(default_tender, "peak_availability_floor", 0.9) * 100
    declared_cuf = _parse_percent_context(normalized, ["annual cuf", "declared cuf", "cuf"], default=getattr(default_tender, "declared_annual_cuf", 0.4) * 100)
    lower_tol = _parse_signed_percent(normalized, ["lower", "below", "minus", "(-)"], default=getattr(default_tender, "cuf_lower_tolerance", 0.15) * 100)
    upper_tol = _parse_signed_percent(normalized, ["upper", "above", "plus", "(+)"], default=getattr(default_tender, "cuf_upper_tolerance", 0.10) * 100)
    multiple = _parse_mw_context(normalized, ["multiple", "minimum block", "bid capacity"], default=getattr(default_tender, "project_mw_multiple", 10.0))
    max_bid = _parse_max_bid(normalized, default=getattr(default_tender, "max_bidder_capacity_mw", tender_size))
    location_allowed = _location_allowed(normalized)
    procurement_type = _procurement_type(normalized)
    storage_required = bool(re.search(r"(?i)\b(energy storage|ESS|BESS|storage system|battery)\b", normalized))
    external_green = bool(re.search(r"(?i)\b(green market|green exchange|external green|other green source|market purchase)\b", normalized))

    settings = {
        "declaredCuf": declared_cuf,
        "hardCompliance": True,
        "externalSupport": external_green,
        "tenderProcurementMw": tender_size,
    }

    tender_title = _tender_title(normalized)
    metadata = _metadata(normalized, name)
    overview = [
        _row("Tender size", _quantum_label(quantum, tender_size), "Defines the procurement quantum; this project's fixed contracted capacity is set in Project Configuration.", chunks, "tender size procurement capacity quantum supply MW MWh"),
        _row("Procurement type", procurement_type, "Determines whether the optimizer must solve firm renewable supply rather than only energy generation.", chunks, "procurement type FDRE firm dispatchable tariff based competitive bidding"),
        _row("Location allowed", location_allowed, "Controls site strategy, ISTS/grid assumptions and whether remote high-resource sites are permitted.", chunks, "location allowed anywhere India ISTS connected renewable energy project"),
        _row("PPA tenor", f"{ppa_years} years", "Sets finance horizon, degradation exposure, BESS augmentation period and DSCR sculpting.", chunks, "PPA term years power purchase agreement"),
        _row("Energy storage", "Required" if storage_required else "Review required", "Storage affects peak availability, dispatchability and capex sizing.", chunks, "energy storage system ESS BESS battery"),
    ]

    technical = [
        _row("Peak-period supply", _peak_obligation_label(quantum, peak_pct, parsed_peak_pct), "Primary sizing driver for BESS MW/MWh and renewable diversity.", chunks, "peak hours peak supply energy requirement contracted capacity MWh kWh"),
        _row("Annual CUF floor", f"{declared_cuf:.1f}% declared annual CUF", "Controls annual PPA energy versus fixed contracted capacity.", chunks, "annual CUF declared generation achieve"),
        _row("CUF tolerance band", f"-{lower_tol:.1f}% / +{upper_tol:.1f}% of declared CUF", "Defines annual lower compliance floor and upper offtake band.", chunks, "CUF within plus minus declared value PPA duration"),
        _row("Bid multiple", f"{multiple:g} MW", "Fixed contracted capacity should be entered in a valid tender multiple.", chunks, "bid capacity multiple minimum block MW"),
        _row("Maximum bidder capacity", f"{max_bid:g} MW", "Caps bidder portfolio exposure in the tender tranche.", chunks, "maximum capacity bidder group affiliates MW"),
        _row("External green support", "Detected" if external_green else "Not clearly detected", "If allowed, limited shortfall energy may be priced instead of forced physical buildout.", chunks, "green market external support shortfall purchase"),
    ]

    constraints = [
        {"Constraint": "Monthly peak availability", "Parsed value": f"{peak_pct:.1f}%", "Model field": "peak_availability_floor", "Confidence": _confidence(peak_pct)},
        {"Constraint": "Declared annual CUF", "Parsed value": f"{declared_cuf:.1f}%", "Model field": "declared_annual_cuf", "Confidence": _confidence(declared_cuf)},
        {"Constraint": "CUF lower tolerance", "Parsed value": f"{lower_tol:.1f}%", "Model field": "cuf_lower_tolerance", "Confidence": _confidence(lower_tol)},
        {"Constraint": "CUF upper tolerance", "Parsed value": f"{upper_tol:.1f}%", "Model field": "cuf_upper_tolerance", "Confidence": _confidence(upper_tol)},
        {"Constraint": "Bid capacity multiple", "Parsed value": f"{multiple:g} MW", "Model field": "project_mw_multiple", "Confidence": _confidence(multiple)},
        {"Constraint": "Maximum bidder capacity", "Parsed value": f"{max_bid:g} MW", "Model field": "max_bidder_capacity_mw", "Confidence": _confidence(max_bid)},
    ]

    commercial = [
        _commercial_row("Tariff objective", "Lowest gross tariff meeting target equity IRR, DSCR and tender compliance.", "Avoids oversizing or pricing to maximize returns instead of clearing target bid economics.", chunks, "tariff quoted bidding tariff competitive bidding"),
        _commercial_row("Penalty / liquidated damages", _extract_condition(chunks, "penalty liquidated damages shortfall generation availability CUF 1.5 tariff") or "Review tender penalty clause", "Penalty clauses decide whether infeasible cases can be priced or must be physically repaired.", chunks, "penalty liquidated damages shortfall 1.5 tariff"),
        _commercial_row("Payment security / PSM", _extract_condition(chunks, "payment security mechanism PSM charge payment") or "Review payment security clauses", "Affects net realized tariff and receivables risk.", chunks, "payment security psm charge"),
        _commercial_row("Bid security", _extract_condition(chunks, "EMD earnest money deposit bid security") or "Review EMD / bid security clause", "Impacts bid working capital and guarantee exposure.", chunks, "EMD earnest money deposit bid security"),
    ]

    financial = _financial_conditions(chunks, text)
    security = _security_formulas(text, tender_size)
    eligibility = _eligibility_rows(chunks, text, tender_size)
    risk_flags = _risk_flags(chunks, text)
    timeline = _timeline_rows(chunks, normalized)
    amendments = _amendment_rows(name, chunks, lower)
    rag_sources = _source_snippets(chunks)
    tender_schema = {
        "title": tender_title,
        "issuer": metadata.get("issuer"),
        "rfs_no": metadata.get("rfs_no"),
        "rfs_date": metadata.get("rfs_date"),
        "search_code": metadata.get("search_code"),
        "procurement_mw": tender_size,
        "peak_supply_mwh": quantum.get("peak_supply_mwh"),
        "peak_hours": quantum.get("peak_hours"),
        "ppa_years": ppa_years,
        "location": location_allowed,
        "storage_required": storage_required,
        "procurement_type": procurement_type,
    }

    result = {
        "source_name": name,
        "tender_schema": tender_schema,
        "security": security,
        "eligibility": eligibility,
        "risk_flags": risk_flags,
        "settings": settings,
        "overview": overview,
        "technical": technical,
        "constraints": constraints,
        "timeline": timeline,
        "commercial": commercial,
        "financial": financial,
        "amendments": amendments,
        "rag_sources": rag_sources,
        "rag_status": {
            "extraction": {**extraction, "pages": len(_pages)},
            "mode": "clause-aware lexical retrieval",
            "chunks": len(chunks),
            "amendment_role": "amendment" if any(k in lower or k in name.lower() for k in AMENDMENT_KEYWORDS) else "base tender",
            "notes": "Haystack-style clause chunking and retrieval are used for structured extraction; OpenAI chat can be layered on the same chunks later.",
        },
    }
    return attach_review(result, text, _pages, payload)


def _parse_tender_size(text: str, default_value: float) -> float:
    patterns = [
        r"(?i)supply\s+of\s+(\d+(?:,\d+)*(?:\.\d+)?)\s*MW",
        r"(?i)(?:selection|procurement).{0,80}?(\d+(?:,\d+)*(?:\.\d+)?)\s*MW",
        r"(?i)(\d+(?:,\d+)*(?:\.\d+)?)\s*MW\s*[\'\"]?firm\s*&?\s*dispatchable",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return _float(match.group(1), default_value)
    return default_value


def _procurement_quantum(text: str) -> dict[str, float | None]:
    result: dict[str, float | None] = {
        "procurement_mw": None,
        "peak_supply_mwh": None,
        "peak_hours": None,
    }
    peak = re.search(
        r"(?i)(\d+(?:,\d+)*(?:\.\d+)?)\s*MWh\s*\(\s*(\d+(?:,\d+)*(?:\.\d+)?)\s*MW\s*[x×]\s*(\d+(?:\.\d+)?)\s*Hrs?",
        text,
    )
    if peak:
        result["peak_supply_mwh"] = _float(peak.group(1), 0.0)
        result["procurement_mw"] = _float(peak.group(2), 0.0)
        result["peak_hours"] = _float(peak.group(3), 0.0)
        return result
    supply = re.search(r"(?i)supply\s+of\s+(\d+(?:,\d+)*(?:\.\d+)?)\s*MW", text)
    if supply:
        result["procurement_mw"] = _float(supply.group(1), 0.0)
    return result


def _quantum_label(quantum: dict[str, float | None], fallback_mw: float) -> str:
    if quantum.get("peak_supply_mwh") and quantum.get("procurement_mw"):
        return f"{quantum['peak_supply_mwh']:g} MWh ({quantum['procurement_mw']:g} MW x {quantum.get('peak_hours') or 4:g} hrs)"
    return f"{fallback_mw:g} MW"


def _peak_obligation_label(quantum: dict[str, float | None], peak_pct: float, parsed_peak_pct: float | None) -> str:
    if quantum.get("peak_supply_mwh") and quantum.get("procurement_mw"):
        per_mw = (float(quantum["peak_supply_mwh"] or 0) * 1000.0) / max(float(quantum["procurement_mw"] or 1), 1.0)
        return f"{per_mw:,.0f} kWh per MW contracted during buyer-selected peak hours"
    if parsed_peak_pct is not None:
        return f"{peak_pct:.1f}% minimum availability during peak hours"
    return f"Default model assumption: {peak_pct:.1f}% minimum peak availability"


def _tender_title(text: str) -> str:
    title_match = re.search(r"(?i)Request\s+for\s+Selection.*?(?=RfS\s+No\.|DISCLAIMER|BID INFORMATION|$)", text)
    if title_match:
        title = normalize_space(title_match.group(0))
        return title[:260]
    first = normalize_space(text[:400])
    return first[:220] if first else "Tender summary"


def _metadata(text: str, name: str) -> dict[str, str]:
    issuer = "Unknown issuer"
    if re.search(r"(?i)Solar Energy Corporation of India|SECI", text[:3000]):
        issuer = "Solar Energy Corporation of India Limited"
    elif re.search(r"(?i)NHPC", text[:3000]):
        issuer = "NHPC Limited"
    rfs = re.search(r"(?i)RfS\s+No\.?\s*[:\-]?\s*([A-Z0-9/&()._\-\s]+?)(?:\s+dated\s+([0-9./-]+)|\s+Page|\s+DISCLAIMER|\s+Tender|$)", text[:8000])
    search = re.search(r"(?i)Tender\s+Search\s+Code[^:\n]{0,100}:\s*([A-Z0-9\-_/]+)", text[:4000])
    return {
        "issuer": issuer,
        "rfs_no": normalize_space(rfs.group(1)) if rfs else name,
        "rfs_date": normalize_space(rfs.group(2)) if rfs and rfs.lastindex and rfs.group(2) else "",
        "search_code": normalize_space(search.group(1)) if search else "",
    }


def _parse_int_context(text: str, terms: list[str], default: int) -> int:
    for term in terms:
        for match in re.finditer(re.escape(term), text, flags=re.IGNORECASE):
            window = text[max(0, match.start() - 120):min(len(text), match.end() + 160)]
            number = re.search(r"(\d{1,2})\s*(?:years?|yrs?)", window, flags=re.IGNORECASE)
            if number:
                return int(number.group(1))
    return default


def _parse_percent_context(text: str, terms: list[str], default: float | None) -> float | None:
    for term in terms:
        for match in re.finditer(re.escape(term), text, flags=re.IGNORECASE):
            # Keep unrelated percentages in adjacent sentences out of this field.
            after = re.split(r"[.;](?:\s|$)", text[match.end():match.end() + 180], maxsplit=1)[0]
            before = re.split(r"[.;](?:\s|$)", text[max(0, match.start() - 160):match.start()])[-1]
            candidates = list(re.finditer(r"(\d+(?:\.\d+)?)\s*%", after))
            candidates += list(re.finditer(r"(\d+(?:\.\d+)?)\s*%", before))[::-1]
            for pct in candidates:
                value = float(pct.group(1))
                if 1 <= value <= 120:
                    return value
    return default


def _parse_signed_percent(text: str, terms: list[str], default: float) -> float:
    for term in terms:
        for match in re.finditer(re.escape(term), text, flags=re.IGNORECASE):
            window = text[max(0, match.start() - 80):min(len(text), match.end() + 120)]
            pct = re.search(r"(\d+(?:\.\d+)?)\s*%", window)
            if pct:
                return float(pct.group(1))
    return default


def _parse_mw_context(text: str, terms: list[str], default: float) -> float:
    for term in terms:
        for match in re.finditer(re.escape(term), text, flags=re.IGNORECASE):
            window = text[max(0, match.start() - 100):min(len(text), match.end() + 120)]
            mw = re.search(r"(\d+(?:,\d+)*(?:\.\d+)?)\s*MW", window, flags=re.IGNORECASE)
            if mw:
                return _float(mw.group(1), default)
    return default


def _parse_max_bid(text: str, default: float) -> float:
    matches = re.finditer(r"(?i)(maximum|max).{0,100}?(\d+(?:,\d+)*(?:\.\d+)?)\s*MW", text)
    for match in matches:
        value = _float(match.group(2), default)
        if value > 0:
            return value
    return default


def _float(value: str, default: float) -> float:
    try:
        return float(value.replace(",", ""))
    except Exception:
        return default


def _location_allowed(text: str) -> str:
    if re.search(r"(?i)anywhere\s+in\s+india", text):
        return "Anywhere in India"
    if re.search(r"(?i)ISTS", text):
        return "ISTS-connected project basis"
    state_match = re.search(r"(?i)\b(Rajasthan|Gujarat|Maharashtra|Karnataka|Tamil Nadu|Andhra Pradesh|Madhya Pradesh)\b", text)
    if state_match:
        return f"Review state/location clause: {state_match.group(1)}"
    return "Review uploaded tender for permitted project location"


def _procurement_type(text: str) -> str:
    if re.search(r"(?i)firm\s*&?\s*dispatchable|FDRE", text):
        return "FDRE / firm and dispatchable renewable power"
    if re.search(r"(?i)round\s*the\s*clock|RTC", text):
        return "RTC renewable power"
    return "Renewable power procurement"


def _row(item: str, detail: str, impact: str, chunks: list[TenderChunk], query: str) -> dict[str, str]:
    source = _best_source(chunks, query)
    return {
        "Item": item,
        "Parsed detail": detail,
        "Bid impact": impact,
        "Source": source,
    }


def _commercial_row(aspect: str, treatment: str, impact: str, chunks: list[TenderChunk], query: str) -> dict[str, str]:
    return {
        "Aspect": aspect,
        "Model treatment": treatment,
        "Bid impact": impact,
        "Source": _best_source(chunks, query),
    }


def _best_source(chunks: list[TenderChunk], query: str) -> str:
    hits = retrieve_chunks(chunks, query, top_k=1)
    if not hits:
        return "No source hit"
    hit = hits[0]
    clause = f", clause {hit.meta.get('clause_id')}" if hit.meta.get("clause_id") else ""
    page = f"p. {hit.meta.get('page_number')}" if hit.meta.get("page_number") else "p. n/a"
    return f"{hit.meta.get('file_name')}, {page}{clause}"


def _extract_condition(chunks: list[TenderChunk], query: str) -> str:
    hits = retrieve_chunks(chunks, query, top_k=1)
    if not hits:
        return ""
    return normalize_space(hits[0].content)[:260]


def _financial_conditions(chunks: list[TenderChunk], text: str) -> list[dict[str, str]]:
    items = [
        ("Document fee", "DOCUMENT FEE COST OF RfS DOCUMENT non refundable amount"),
        ("Bid processing fee", "BID PROCESSING FEE quoted capacity maximum amount"),
        ("EMD / earnest money", "Earnest Money Deposit EMD formula installed capacity solar wind ESS"),
        ("PBG / performance security", "Performance Bank Guarantee PBG formula installed capacity solar wind ESS"),
        ("Success charge", "success charge trading margin charge"),
        ("Payment security / PSM", "payment security mechanism PSM charge payment security fund"),
        ("Penalty / liquidated damages", "penalty liquidated damages shortfall availability CUF 1.5 times tariff"),
        ("Net worth / financial eligibility", "net worth financial eligibility turnover line of credit"),
        ("Tariff and payment", "tariff quoted kWh payment energy charges tariff adoption"),
    ]
    rows = []
    for item, query in items:
        direct = _direct_financial_snippet(item, text)
        hits = retrieve_chunks(chunks, query, top_k=1)
        detail = direct or (normalize_space(hits[0].content)[:360] if hits else "Not parsed from upload")
        rows.append({
            "Financial condition": item,
            "Extracted clause detail": detail,
            "Bid impact": _financial_impact(item),
            "Source": _direct_financial_source(item) if direct else (_source_label(hits[0]) if hits else "No source hit"),
        })
    return rows


def _direct_financial_snippet(item: str, text: str) -> str:
    patterns = {
        "Document fee": r"(?is)DOCUMENT\s+FEE.*?(?:response\s+to\s+RfS|APPLICABLE)",
        "Bid processing fee": r"(?is)BID\s+PROCESSING\s+FEE.*?(?:response\s+to\s+RfS|APPLICABLE)",
        "EMD / earnest money": r"(?is)\b16\s+Earnest\s+Money\s+Deposit\s*\(EMD\).*?(?=\n\s*16\.2|\n\s*17\s+Performance|\f|$)",
        "PBG / performance security": r"(?is)\b17\s+Performance\s+Bank\s+Guarantee.*?(?=\n\s*17\.2|\n\s*18\s+|\f|$)",
        "Net worth / financial eligibility": r"(?is)(?:minimum\s+annual\s+turnover|net\s+worth).*?(?:bid\s+submission\s+deadline|PBDIT\s+certificate).*?(?=\n\s*37|\f|$)",
    }
    pattern = patterns.get(item)
    if not pattern:
        return ""
    match = re.search(pattern, text)
    if not match:
        return ""
    return normalize_space(match.group(0))[:520]


def _direct_financial_source(item: str) -> str:
    sources = {
        "Document fee": "Bid information sheet",
        "Bid processing fee": "Bid information sheet",
        "EMD / earnest money": "Clause 16",
        "PBG / performance security": "Clause 17.1",
        "Net worth / financial eligibility": "Financial eligibility clauses",
    }
    return sources.get(item, "Tender clause")


def _financial_impact(item: str) -> str:
    impact = {
        "Document fee": "Non-refundable bid participation cost.",
        "Bid processing fee": "Bid transaction cost linked to quoted capacity.",
        "EMD / earnest money": "Bid-stage guarantee/cash exposure.",
        "PBG / performance security": "Post-award bank guarantee exposure and cost.",
        "Processing fee": "Bid transaction cost.",
        "Success charge": "Award cost included in project capex/fees.",
        "Payment security / PSM": "Receivable risk and net tariff adjustment.",
        "Penalty / liquidated damages": "Defines cost of non-compliance and hard/soft constraint treatment.",
        "Net worth / financial eligibility": "Bidder qualification and consortium structuring.",
        "Tariff and payment": "Revenue model and bid competitiveness.",
    }
    return impact.get(item, "Review impact before bid submission.")


def _money_per_unit(text: str, label: str, unit: str) -> float | None:
    pattern = rf"(?i)INR\s+([\d,]+)\s*x\s*Rated[^.\n]{{0,140}}{label}[^.\n]{{0,80}}\({unit}\)"
    match = re.search(pattern, text)
    if not match:
        return None
    return _float(match.group(1), 0.0)


def _security_formulas(text: str, procurement_mw: float) -> list[dict[str, str]]:
    emd_solar = emd_wind = emd_ess = None
    emd_match = re.search(r"(?is)Earnest\s+Money\s+Deposit\s*=\s*\[(.*?)\]\.", text)
    if emd_match:
        rates = _formula_rates(emd_match.group(1))
        emd_solar, emd_wind, emd_ess = rates.get("solar"), rates.get("wind"), rates.get("ess")
    pbg_solar = pbg_wind = pbg_ess = None
    pbg_match = re.search(r"(?is)Performance\s+Bank\s+Guarantee\s*=\s*\[(.*?)\]\.", text)
    if pbg_match:
        rates = _formula_rates(pbg_match.group(1))
        pbg_solar, pbg_wind, pbg_ess = rates.get("solar"), rates.get("wind"), rates.get("ess")
    rows = []
    rows.append({
        "Instrument": "EMD",
        "Formula": _formula_label(emd_solar, emd_wind, emd_ess),
        "Validity / timing": _snippet_or_default(text, r"(?is)valid\s+for\s+12\s+months.*?claim\s+period.*?30\s+days", "BG / POI / Surety; validity and claim period as per tender."),
        "Clause": "16",
    })
    rows.append({
        "Instrument": "PBG",
        "Formula": _formula_label(pbg_solar, pbg_wind, pbg_ess),
        "Validity / timing": _snippet_or_default(text, r"(?is)prior\s+to\s+signing\s+of\s+PPA.*?9\s+months\s+after\s+the\s+SCSD", "Submitted before PPA signing; validity as per tender."),
        "Clause": "17.1",
    })
    processing = re.search(r"(?is)BID\s+PROCESSING\s+FEE\s*(.*?)(?:APPLICABLE|NOT\s+APPLICABLE)", text)
    rows.append({
        "Instrument": "Bid processing fee",
        "Formula": normalize_space(processing.group(1))[:180] if processing else "Review bid information sheet",
        "Validity / timing": "Submitted with response to RfS",
        "Clause": "Bid information sheet",
    })
    document_fee = re.search(r"(?is)DOCUMENT\s+FEE.*?Amount:\s*(.*?)(?:APPLICABLE|NOT\s+APPLICABLE)", text)
    rows.append({
        "Instrument": "Document fee",
        "Formula": normalize_space(document_fee.group(1))[:180] if document_fee else "Review bid information sheet",
        "Validity / timing": "Non-refundable; submitted with response to RfS",
        "Clause": "Bid information sheet",
    })
    return rows


def _formula_rates(block: str) -> dict[str, float]:
    rates: dict[str, float] = {}
    pieces = re.split(r"\+\s*", block)
    for piece in pieces:
        amount = re.search(r"(?i)INR\s+([\d,]+)", piece)
        if not amount:
            continue
        value = _float(amount.group(1), 0.0)
        lowered = piece.lower()
        if "solar" in lowered:
            rates["solar"] = value
        elif "wind" in lowered or "other re" in lowered:
            rates["wind"] = value
        elif "ess" in lowered or "storage" in lowered:
            rates["ess"] = value
    return rates


def _formula_label(solar: float | None, wind: float | None, ess: float | None) -> str:
    parts = []
    if solar:
        parts.append(f"Rs {solar / 100000:.2f} L/MW solar")
    if wind:
        parts.append(f"Rs {wind / 100000:.2f} L/MW wind/other RE")
    if ess:
        parts.append(f"Rs {ess / 100000:.2f} L/MWh ESS")
    return " + ".join(parts) if parts else "Formula not confidently parsed"


def _snippet_or_default(text: str, pattern: str, default: str) -> str:
    match = re.search(pattern, text)
    return normalize_space(match.group(0))[:180] if match else default


def _eligibility_rows(chunks: list[TenderChunk], text: str, procurement_mw: float) -> list[dict[str, str]]:
    turnover_match = re.search(r"(?i)annual\s+turnover\s+of\s+INR\s+([\d,]+)\s*/\s*MW", text)
    pbdit_match = re.search(r"(?i)PBDIT.*?INR\s+([\d,]+)\s*/\s*MW", text)
    loc_match = re.search(r"(?i)Line\s+of\s+Credit.*?INR\s+([\d,]+)\s*/\s*MW", text)
    rows = [
        _eligibility_row("Annual turnover", turnover_match, procurement_mw, "Excludes other income; CA/statutory auditor evidence required.", "36.2"),
        _eligibility_row("PBDIT", pbdit_match, procurement_mw, "Positive internal resource generation, excluding other and exceptional income.", "36.2"),
        _eligibility_row("Line of credit", loc_match, procurement_mw, "Working-capital support letter from lender/bank.", "36.2"),
    ]
    mse = _extract_condition(chunks, "MSE exemption document fee bid processing fee EMD UDYAM consortium")
    if mse:
        rows.append({
            "Requirement": "MSE exemption",
            "Threshold": "Fees and EMD exemption if eligible",
            "Bidder action": "Verify valid UDYAM registration and consortium-wide eligibility.",
            "Clause": "15",
            "Evidence": mse[:220],
        })
    return rows


def _eligibility_row(name: str, match: re.Match | None, procurement_mw: float, action: str, clause: str) -> dict[str, str]:
    if match:
        per_mw = _float(match.group(1), 0.0)
        threshold_cr = per_mw * procurement_mw / 10000000.0
        threshold = f"Rs {per_mw / 10000000.0:.4f} Cr/MW; Rs {threshold_cr:,.1f} Cr for {procurement_mw:g} MW"
        evidence = normalize_space(match.group(0))[:220]
    else:
        threshold = "Not confidently parsed"
        evidence = "Review financial eligibility clause"
    return {
        "Requirement": name,
        "Threshold": threshold,
        "Bidder action": action,
        "Clause": clause,
        "Evidence": evidence,
    }


def _risk_flags(chunks: list[TenderChunk], text: str) -> list[dict[str, str]]:
    candidates = [
        ("Critical", "No deviations permitted", "deviations not permitted exceptions rejected without opening partial EMD BG format"),
        ("High", "Connectivity / GNA timing", "GNA connectivity scheduled commencement supply SCSD delay charges"),
        ("Medium", "Payment counterparty / buying entity", "buying entities unnamed payment security fund receivable risk"),
        ("Medium", "BG format and offline originals", "Bank Guarantee Format original offline pass phrase submitted"),
        ("Opportunity", "Merchant use outside peak hours", "outside the Peak Hours sell power exchange without NOC"),
    ]
    rows = []
    for severity, title, query in candidates:
        hit = retrieve_chunks(chunks, query, top_k=1)
        if not hit:
            continue
        rows.append({
            "Severity": severity,
            "Risk / opportunity": title,
            "Bid implication": _risk_implication(title),
            "Source": _source_label(hit[0]),
            "Evidence": normalize_space(hit[0].content)[:260],
        })
    return rows


def _risk_implication(title: str) -> str:
    mapping = {
        "No deviations permitted": "Any exception, format deviation or incomplete security can reject the bid.",
        "Connectivity / GNA timing": "SCSD and GNA queue risk can become delay charges or PBG exposure.",
        "Payment counterparty / buying entity": "Counterparty/payment-security terms affect receivable risk and tariff buffer.",
        "BG format and offline originals": "Treasury/legal must prepare exact instruments and offline submissions on time.",
        "Merchant use outside peak hours": "Idle RE/BESS can create upside if exchange sales are allowed.",
    }
    return mapping.get(title, "Review before bid submission.")


def _timeline_rows(chunks: list[TenderChunk], text: str) -> list[dict[str, str]]:
    critical_dates = _critical_date_map(text)
    milestones = [
        ("Tender issue / RfS date", "RFS date issued tender issue date"),
        ("Pre-bid / clarification", "pre bid meeting clarification last date query"),
        ("Bid submission deadline", "bid submission last date due date online submission"),
        ("E-reverse auction", "e reverse auction e-RA reverse auction"),
        ("PPA execution", "PPA execution signing power purchase agreement"),
        ("Financial closure", "financial closure months from PPA"),
        ("Scheduled commissioning", "scheduled commissioning SCOD commissioning months"),
    ]
    rows = []
    for milestone, query in milestones:
        direct = critical_dates.get(milestone)
        hits = retrieve_chunks(chunks, query, top_k=3)
        detail = direct or _timeline_detail(hits) or "Not parsed from upload"
        rows.append({
            "Milestone": milestone,
            "Parsed detail": detail,
            "Bid impact": "Impacts bid validity, EPC quote validity, financing timetable and execution risk.",
            "Source": "Critical dates table" if direct else (_source_label(hits[0]) if hits else "No source hit"),
        })
    return rows


def _critical_date_map(text: str) -> dict[str, str]:
    labels = [
        ("Tender issue / RfS date", r"(?i)(?:publishing|issue|rfs)\s+date\s*&?\s*time"),
        ("Pre-bid / clarification", r"(?i)pre\s*bid\s*meeting\s+date\s*&?\s*time"),
        ("Bid submission deadline", r"(?i)online\s+bid\s+submission\s+closing\s+date\s*&?\s*time"),
        ("E-reverse auction", r"(?i)(?:start\s+of\s+)?e[-\s]*reverse\s+auction"),
        ("Technical bid opening", r"(?i)technical\s+bid.*?(?:opening|submission)"),
        ("Offline submission", r"(?i)last\s+date\s+of\s+offline\s+submission"),
        ("Bid submission start", r"(?i)online\s+bid\s+submission\s+start\s+date\s*&?\s*time"),
        ("Clarification deadline", r"(?i)last\s+date\s+of\s+receipt\s+of\s+queries\s*/?clarification"),
    ]
    mapped: dict[str, str] = {}
    for milestone, pattern in labels:
        match = re.search(pattern, text)
        if not match:
            continue
        window = text[match.start():min(len(text), match.end() + 260)]
        date = DATE_RE.search(window)
        if date:
            mapped[milestone] = normalize_space(window[:date.end()])
        elif re.search(r"(?i)intimated\s+separately|to\s+be\s+intimated", window):
            mapped[milestone] = normalize_space(window[:180])
    if "Pre-bid / clarification" not in mapped and "Clarification deadline" in mapped:
        mapped["Pre-bid / clarification"] = mapped["Clarification deadline"]
    return mapped


def _timeline_detail(hits: list[TenderChunk]) -> str:
    for hit in hits:
        content = normalize_space(hit.content)
        date = DATE_RE.search(content)
        if date:
            start = max(0, date.start() - 90)
            end = min(len(content), date.end() + 90)
            return content[start:end]
        months = re.search(r"(?i)\b\d+\s*months?\b.{0,80}(?:PPA|commissioning|SCOD|effective date|LOA)", content)
        if months:
            return months.group(0)
    return ""


def _amendment_rows(name: str, chunks: list[TenderChunk], lower: str) -> list[dict[str, str]]:
    is_amendment = document_role(name, " ".join(chunk.content for chunk in chunks)) == "amendment"
    rows: list[dict[str, str]] = []
    if not is_amendment:
        return [{
            "Document role": "Base tender",
            "Clause reference": "N/A",
            "Change type": "No amendment detected",
            "Impact summary": "Upload corrigendum/addendum files to track clause-level changes against the base tender.",
            "Source": name,
        }]
    for hit in retrieve_chunks(chunks, "amendment corrigendum addendum clarification modified clause", top_k=12):
        refs = CHANGE_REF_RE.findall(hit.content)
        rows.append({
            "Document role": "Amendment / corrigendum",
            "Clause reference": ", ".join(refs[:4]) if refs else hit.meta.get("clause_id", "Review text"),
            "Change type": "Potential modification",
            "Impact summary": normalize_space(hit.content)[:300],
            "Source": _source_label(hit),
        })
    return rows[:10] or [{
        "Document role": "Amendment / corrigendum",
        "Clause reference": "Review text",
        "Change type": "Potential modification",
        "Impact summary": "Amendment keywords detected, but no clause reference was confidently extracted.",
        "Source": name,
    }]


def _source_snippets(chunks: list[TenderChunk]) -> list[dict[str, str]]:
    themes = [
        ("Tender summary", "tender size procurement capacity location ppa"),
        ("Key dates", "bid submission date pre bid reverse auction commissioning"),
        ("Financial conditions", "EMD PBG processing fee success charge penalty payment security"),
        ("Technical compliance", "peak availability annual CUF energy storage contracted capacity"),
        ("Amendments", "amendment corrigendum addendum clarification clause"),
    ]
    snippets = []
    seen = set()
    for theme, query in themes:
        for hit in retrieve_chunks(chunks, query, top_k=2):
            key = (theme, hit.meta.get("page_number"), hit.meta.get("split_id"))
            if key in seen:
                continue
            seen.add(key)
            snippets.append({
                "Theme": theme,
                "Source": _source_label(hit),
                "Snippet": normalize_space(hit.content)[:420],
                "Score": f"{hit.score:.3f}",
            })
    return snippets[:12]


def _source_label(hit: TenderChunk) -> str:
    clause = f", clause {hit.meta.get('clause_id')}" if hit.meta.get("clause_id") else ""
    return f"{hit.meta.get('file_name')}, p. {hit.meta.get('page_number')}{clause}"


def _confidence(value: Any) -> str:
    if value is None:
        return "Needs review"
    try:
        number = float(value)
    except Exception:
        return "Medium"
    return "High" if number > 0 else "Needs review"
