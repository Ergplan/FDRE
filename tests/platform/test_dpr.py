"""Day 6 gate: DPR values trace to their approved source run. The award freezes the run's
baseline; traced values cannot be edited; narrative is versioned with revision checks;
exports are stored artifacts linked to the run's workbook."""

from __future__ import annotations

import copy
import importlib.util
import io
import pathlib
import shutil

import pytest
from docx import Document
from sqlalchemy import text
from sqlalchemy.exc import DBAPIError

_spec = importlib.util.spec_from_file_location("econ_helpers2", pathlib.Path(__file__).with_name("test_economic_workflows.py"))
_econ = importlib.util.module_from_spec(_spec)
assert _spec.loader is not None
_spec.loader.exec_module(_econ)


def test_award_freezes_baseline_and_dpr_values_trace_to_the_source_run(client, seed, as_user, worker, session_factory):
    org, bidder, scenario = _econ._scenario(client, seed, as_user)
    run = _econ._run(client, worker, org, bidder, scenario["id"], "fixed_tariff", {"tariff": 4.2, "dispatch": "milp"})
    assert run["status"] == "succeeded", run.get("error")
    tid = scenario["tender_id"]
    award_body = {"run_id": run["id"], "tariff_inr_per_kwh": 4.2, "capacity_mw": 50, "award_date": "2026-10-15", "letter_reference": "LoA/2026/17"}
    # recording an award needs bid.approve
    assert client.post(f"/api/v1/orgs/{org}/tenders/{tid}/awards", json=award_body, headers=bidder).status_code == 403
    assert (
        client.post(
            f"/api/v1/orgs/{org}/users", json={"email": "appr@acme.test", "display_name": "A", "password": "correct-horse", "role": "approver"}, headers=bidder
        ).status_code
        == 201
    )
    client.post(f"/api/v1/orgs/{org}/tenders/{tid}/members", json={"email": "appr@acme.test"}, headers=bidder)
    approver = as_user("appr@acme.test")
    award = client.post(f"/api/v1/orgs/{org}/tenders/{tid}/awards", json=award_body, headers=approver)
    assert award.status_code == 201, award.text
    award = award.json()
    assert award["baseline"]["run"]["inputs_hash"] == run["snapshot"]["inputs"]["hash"]
    assert award["baseline"]["result"]["finance"]["equity_irr"] == run["result"]["finance"]["equity_irr"]
    assert award["baseline"]["inputs"]["supply"]["modules"][0]["name"] == "Solar A" and award["status"] == "recorded"
    # the scenario moves on: the run goes stale, the baseline does not move, and a stale run cannot be awarded
    new_inputs = {**_econ._inputs(), "settings": {"changed": True}}
    assert (
        client.put(
            f"/api/v1/orgs/{org}/scenarios/{scenario['id']}/inputs", json={"expected_revision": scenario["revision"], "inputs": new_inputs}, headers=bidder
        ).status_code
        == 200
    )
    assert client.get(f"/api/v1/orgs/{org}/runs/{run['id']}?include_result=false", headers=bidder).json()["stale"] is True
    stale = client.post(f"/api/v1/orgs/{org}/tenders/{tid}/awards", json=award_body, headers=approver)
    assert stale.status_code == 422 and "stale" in stale.json()["message"]
    assert client.get(f"/api/v1/orgs/{org}/awards/{award['id']}", headers=bidder).json()["baseline"]["inputs"]["settings"] == {}
    with session_factory() as s, pytest.raises(DBAPIError, match="frozen"):
        s.execute(text("update award set awarded_tariff_inr_per_kwh = 9.9 where id = :id"), {"id": award["id"]})
        s.commit()
    # the DPR workspace is generated from the baseline; every value names its run and path
    ws = client.post(f"/api/v1/orgs/{org}/awards/{award['id']}/dpr", headers=bidder)
    assert ws.status_code == 201, ws.text
    ws = ws.json()
    sections = ws["version"]["sections"]
    assert [s["key"] for s in sections] == [
        "project_overview",
        "technical_configuration",
        "energy_and_dispatch",
        "financial_model",
        "tender_obligations",
        "assumptions_and_sources",
        "risks_and_limitations",
        "annexures",
    ]
    fin = next(s for s in sections if s["key"] == "financial_model")
    irr = next(v for v in fin["values"] if v["key"] == "equity_irr")
    assert irr["value"] == run["result"]["finance"]["equity_irr"] and irr["source"] == {"run_id": run["id"], "path": "result.finance.equity_irr"}
    tech = next(s for s in sections if s["key"] == "technical_configuration")
    assert any(v["path"] if False else v["source"]["path"] == "inputs.supply.modules[0].mwac" and v["value"] == 120.0 for v in tech["values"])
    trace = client.get(f"/api/v1/orgs/{org}/dpr/{ws['id']}/trace", headers=bidder).json()
    assert trace["all_match"] is True and trace["run_id"] == run["id"] and len(trace["values"]) > 15 and trace["run_stale"] is True
    assert all(v["current"] == v["value"] for v in trace["values"] if v["key"] in ("equity_irr", "awarded_tariff", "ppa_mwh_y1"))
    # traced values cannot be edited; narrative can, as a new version with a revision check
    tampered = copy.deepcopy(sections)
    next(s for s in tampered if s["key"] == "financial_model")["values"][0]["value"] = 0.99
    r = client.put(
        f"/api/v1/orgs/{org}/dpr/{ws['id']}", json={"expected_revision": ws["revision"], "sections": tampered, "change_note": "tamper"}, headers=bidder
    )
    assert r.status_code == 422 and "cannot be edited" in r.json()["message"]
    edited = copy.deepcopy(sections)
    edited[0]["body"] = "The project is a 50 MW FDRE supply to NHPC from a solar plus storage plant."
    v2 = client.put(
        f"/api/v1/orgs/{org}/dpr/{ws['id']}",
        json={"expected_revision": ws["revision"], "sections": edited, "change_note": "overview narrative"},
        headers=bidder,
    )
    assert v2.status_code == 200, v2.text
    v2 = v2.json()
    assert v2["current_version_no"] == 2 and v2["revision"] == 2 and v2["version"]["changed_sections"] == ["project_overview"]
    assert next(v for v in v2["version"]["sections"][3]["values"] if v["key"] == "equity_irr")["value"] == irr["value"]
    assert client.put(f"/api/v1/orgs/{org}/dpr/{ws['id']}", json={"expected_revision": 1, "sections": edited}, headers=bidder).status_code == 409
    client.post(f"/api/v1/orgs/{org}/tenders/{tid}/members", json={"email": "viewer@acme.test"}, headers=bidder)
    assert (
        client.put(f"/api/v1/orgs/{org}/dpr/{ws['id']}", json={"expected_revision": 2, "sections": edited}, headers=as_user("viewer@acme.test")).status_code
        == 403
    )
    assert client.get(f"/api/v1/orgs/{org}/dpr/{ws['id']}", headers=as_user("admin2@other.test")).status_code == 404
    assert [v["version_no"] for v in client.get(f"/api/v1/orgs/{org}/dpr/{ws['id']}/versions", headers=bidder).json()] == [1, 2]
    with session_factory() as s, pytest.raises(DBAPIError, match="immutable"):
        s.execute(text("update dpr_version set change_note = 'x' where workspace_id = :id"), {"id": ws["id"]})
        s.commit()
    # exports: DOCX always; PDF through LibreOffice; linked to the run's workbook once it exists
    docx = client.post(f"/api/v1/orgs/{org}/dpr/{ws['id']}/export?format=docx", headers=bidder)
    assert docx.status_code == 201, docx.text
    docx = docx.json()
    assert docx["kind"] == "docx" and docx["version_no"] == 2 and docx["linked_workbook_sha256"] is None
    blob = client.get(docx["download"], headers=bidder)
    assert blob.status_code == 200 and blob.headers["X-Content-SHA256"] == docx["sha256"]
    textual = (
        "\n".join(p.text for p in Document(io.BytesIO(blob.content)).paragraphs)
        + "\n"
        + "\n".join(c.text for t in Document(io.BytesIO(blob.content)).tables for row in t.rows for c in row.cells)
    )
    assert "4.2 INR/kWh" in textual and run["id"] in textual and "result.finance.equity_irr" in textual and "LoA/2026/17" in textual
    wb = client.post(f"/api/v1/orgs/{org}/runs/{run['id']}/workbook?verify=false", headers=bidder).json()
    docx2 = client.post(f"/api/v1/orgs/{org}/dpr/{ws['id']}/export?format=docx", headers=bidder).json()
    assert docx2["linked_workbook_sha256"] == wb["artifact"]["sha256"]
    if shutil.which("soffice"):
        pdf = client.post(f"/api/v1/orgs/{org}/dpr/{ws['id']}/export?format=pdf", headers=bidder).json()
        assert pdf["kind"] == "pdf" and client.get(pdf["download"], headers=bidder).content.startswith(b"%PDF")
    kinds = [a["kind"] for a in client.get(f"/api/v1/orgs/{org}/dpr/{ws['id']}/artifacts", headers=bidder).json()]
    assert kinds.count("docx") == 2
    # issuing is an approval; an issued DPR is final
    assert client.post(f"/api/v1/orgs/{org}/dpr/{ws['id']}/status", json={"expected_revision": 2, "status": "issued"}, headers=bidder).status_code == 403
    issued = client.post(f"/api/v1/orgs/{org}/dpr/{ws['id']}/status", json={"expected_revision": 2, "status": "issued"}, headers=approver).json()
    assert issued["status"] == "issued" and issued["revision"] == 3
    assert client.put(f"/api/v1/orgs/{org}/dpr/{ws['id']}", json={"expected_revision": 3, "sections": edited}, headers=bidder).status_code == 422
    events = client.get(f"/api/v1/orgs/{org}/audit?object_type=dpr_workspace&object_id={ws['id']}", headers=bidder).json()
    assert {"dpr.create", "dpr.version", "dpr.export", "dpr.status"} <= {e["action"] for e in events}
