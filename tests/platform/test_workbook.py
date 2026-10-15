"""Day 5 gate: the formula workbook recalculated by a real spreadsheet engine (LibreOffice
headless) reconciles with finance.v1, and changing an input changes dependent outputs."""

from __future__ import annotations

import io
import shutil

import pytest
from openpyxl import load_workbook

from fdre_platform.finance import workbook
from fdre_platform.finance.model import CapexItem, FinanceAssumptions, OperatingYear, OpexItem, evaluate

pytestmark = pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice not available: workbook verification blocked")


def _case():
    a = FinanceAssumptions(
        ppa_years=12,
        construction_months=24,
        drawdown_profile=[0.45, 0.55],
        capex=[CapexItem("Solar EPC", 3.1e9, source="test"), CapexItem("Land", 2.0e8, "land_non_depreciable", source="test")],
        contingency_fraction=0.03,
        debt_fraction=0.7,
        interest_rate=0.09,
        tenor_years=10,
        repayment="sculpted",
        sculpt_target_dscr=1.2,
        upfront_fee_fraction=0.01,
        dsra_months=6,
        receivable_days=45,
        working_capital_interest_rate=0.1,
        tax_rate=0.25168,
        tax_wdv_rate=0.4,
        tax_wdv_eligible_fraction=0.95,
        book_depreciation_years=25,
        residual_value_fraction=0.05,
        discount_rate=0.1,
        target_equity_irr=0.14,
        min_dscr=1.1,
        opex=[OpexItem("O&M", 4.0e7, 0.035, source="test")],
        augmentation=[(6, 1.5e8)],
        sources={"finance_block": "test fixture"},
    )
    ops = [
        OperatingYear(ppa_mwh=220000.0 * (0.995**i), merchant_mwh=5000.0, merchant_revenue_inr=1.5e7, penalty_mwh=100.0 if i == 3 else 0.0) for i in range(12)
    ]
    return a, ops


def test_workbook_recalculates_and_reconciles_with_backend():
    a, ops = _case()
    tariff = 3.4
    res = evaluate(a, ops, tariff, psm_charge_inr_per_kwh=0.02, penalty_multiplier=1.5)
    assert res.equity_irr is not None
    data = workbook.build(
        a,
        ops,
        res,
        provenance={"run_id": "test", "inputs_hash": "abc", "rules_version": "rule_set:1", "solver": "test", "sensitivities": []},
        tariff=tariff,
        psm=0.02,
        penalty_multiplier=1.5,
    )
    wb = load_workbook(io.BytesIO(data))
    assert {"Read Me", "Inputs", "Dispatch", "Capex & Funding", "Operations", "Debt Tax Cashflow", "Returns", "Sensitivities", "Checks"} <= set(wb.sheetnames)
    # Real formulas, no external links.
    formulas = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) > 300 and not any("[" in f for f in formulas)
    assert "tariff" in wb.defined_names and "equity_irr" in wb.defined_names
    verdict = workbook.verify(data, res)
    assert verdict["status"] == "verified", verdict
    by = {c["metric"]: c for c in verdict["comparisons"]}
    assert by["equity_irr"]["ok"] and by["npv_equity"]["ok"] and by["min_dscr"]["ok"] and by["project_irr"]["ok"]
    assert all(s == "OK" for s in verdict["checks_sheet"]), verdict["checks_sheet"]


def test_changing_an_input_changes_dependent_outputs():
    a, ops = _case()
    res = evaluate(a, ops, 3.4, psm_charge_inr_per_kwh=0.02, penalty_multiplier=1.5)
    data = workbook.build(a, ops, res, provenance={"run_id": "t", "sensitivities": []}, tariff=3.4, psm=0.02, penalty_multiplier=1.5)
    base = workbook.recalculate_with_libreoffice(data)["values"]["Returns"]["B5"]
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Inputs"]
    row = next(r for r in range(2, 40) if ws.cell(r, 1).value == "Tariff")
    ws.cell(row, 2).value = 3.4 * 1.1
    out = io.BytesIO()
    wb.save(out)
    changed = workbook.recalculate_with_libreoffice(out.getvalue())["values"]
    assert changed["Returns"]["B5"] > base + 0.005  # higher tariff, higher equity IRR
    assert changed["Checks"]["E4"] != "OK" and changed["Checks"]["E8"] != "OK"  # revenue and IRR checks report the divergence from the exported values
    # And the backend agrees with the recalculated workbook at the new tariff.
    res2 = evaluate(a, ops, 3.4 * 1.1, psm_charge_inr_per_kwh=0.02, penalty_multiplier=1.5)
    assert changed["Returns"]["B5"] == pytest.approx(res2.equity_irr, abs=2e-3)  # stated simplifications (Read Me)
