"""Formula-based Excel model mirroring finance.v1.

Every computed cell is an Excel formula referring to the Inputs and Dispatch sheets; the
only hard-coded numbers are inputs (blue) and the imported dispatch snapshot (grey, marked
as a fixed run snapshot). A Checks sheet reconciles the workbook against the backend values
stored next to each result so a reader sees where the two agree, and the Read Me says what
recalculates in Excel (finance) and what does not (dispatch, which needs a new run).

Simplifications against finance.v1, stated on the Read Me sheet: the DSCR-sculpted
repayment is computed with the same per-year rule (principal = CFADS/target - interest,
remainder in the last tenor year) using the interest of the same year on the opening
balance; the tax/interest/CFADS circularity is broken by computing tax on EBITDA less
interest on the opening balance (as finance.v1 does after convergence); IDC and fees are
taken from the backend's converged construction schedule as labelled inputs; the DSRA is a
formula on the first year's debt service.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from .model import FinanceAssumptions, FinanceResult, OperatingYear

INPUT_FONT = Font(color="0000FF")
INPUT_FILL = PatternFill("solid", fgColor="FFF8DC")
FORMULA_FONT = Font(color="000000")
SNAPSHOT_FILL = PatternFill("solid", fgColor="EEEEEE")
HEADER_FONT = Font(bold=True)
CHECK_FILL = PatternFill("solid", fgColor="E8F5E9")
WORKBOOK_VERSION = "workbook.v1 (mirrors finance.v1)"


def _name(wb: Workbook, name: str, ref: str) -> None:
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def build(assumptions: FinanceAssumptions, operating: list[OperatingYear], result: FinanceResult, *, provenance: dict[str, Any], tariff: float, psm: float, penalty_multiplier: float) -> bytes:
    a = assumptions
    n = a.ppa_years
    wb = Workbook()
    # ---------------- Read Me
    ws = wb.active
    ws.title = "Read Me"
    rows = [
        ["FDRE enterprise platform: financial model workbook", ""],
        ["Workbook version", WORKBOOK_VERSION],
        ["Generated at (UTC)", datetime.now(UTC).isoformat()],
        ["Run id", provenance.get("run_id")],
        ["Scenario", provenance.get("scenario_name")],
        ["Tender", provenance.get("tender_name")],
        ["Inputs hash", provenance.get("inputs_hash")],
        ["Rules version", provenance.get("rules_version")],
        ["Solver / formulation", provenance.get("solver")],
        ["Finance engine", a.version],
        ["", ""],
        ["How to use", "Blue cells on Inputs are editable assumptions. Every other number is an Excel formula. Changing an input recalculates the financial sheets (revenue, debt, tax, cash flows, IRR, NPV, DSCR)."],
        ["What does NOT recalculate", "The Dispatch sheet is a fixed snapshot of the optimisation run named above (grey cells). Changing capacities, profiles, losses, tender rules or any dispatch-affecting input requires a new application run; the workbook cannot re-run the MILP."],
        ["Checks", "The Checks sheet compares workbook results with the backend values recorded at export time. They agree at export; after you change inputs they are expected to differ and the check says so."],
        ["Simplifications", "DSCR-sculpted principal = CFADS/target DSCR - interest on the opening balance; tax uses interest on the opening balance; IDC and fees are taken from the converged backend construction schedule and shown as inputs (labelled); the DSRA is months of year-1 debt service. After an input change the workbook and a fresh backend run agree to within a few tenths of a percentage point of IRR because of these simplifications."],
        ["No external links", "This workbook has no links to other files and no hidden circular references; iterative calculation is not required."],
    ]
    for r in rows:
        ws.append(r)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 120
    ws["A1"].font = Font(bold=True, size=13)
    # ---------------- Inputs
    inp = wb.create_sheet("Inputs")
    inp.append(["Assumption", "Value", "Unit", "Source"])
    for c in inp[1]:
        c.font = HEADER_FONT
    scalar: list[tuple[str, str, Any, str, str]] = [
        ("tariff", "Tariff", tariff, "INR/kWh", provenance.get("tariff_source", "run")),
        ("psm", "PSM / trading charge deducted", psm, "INR/kWh", "tender rules"),
        ("penalty_mult", "Penalty multiple on shortfall", penalty_multiplier, "x tariff", "tender rules"),
        ("ppa_years", "PPA years", n, "years", "tender rules"),
        ("contingency", "Contingency", a.contingency_fraction, "fraction of hard cost", a.sources.get("finance_block", "defaulted")),
        ("debt_fraction", "Debt fraction", a.debt_fraction, "fraction of project cost", a.sources.get("finance_block", "defaulted")),
        ("interest", "Interest rate", a.interest_rate, "per year", a.sources.get("finance_block", "defaulted")),
        ("tenor", "Debt tenor", min(a.tenor_years, n), "years", a.sources.get("finance_block", "defaulted")),
        ("target_dscr", "Sculpting target DSCR", a.sculpt_target_dscr, "x", a.sources.get("finance_block", "defaulted")),
        ("idc", "IDC (converged backend construction schedule)", result.idc_inr, "INR", "backend finance.v1 (labelled input)"),
        ("fees", "Upfront fees", result.fees_inr, "INR", "backend finance.v1 (labelled input)"),
        ("dsra_months", "DSRA (months of debt service)", a.dsra_months, "months", a.sources.get("finance_block", "defaulted")),
        ("receivable_days", "Receivable days", a.receivable_days, "days", a.sources.get("finance_block", "defaulted")),
        ("wc_rate", "Working capital interest", a.working_capital_interest_rate, "per year", a.sources.get("finance_block", "defaulted")),
        ("tax_rate", "Corporate tax rate", a.tax_rate, "fraction", a.sources.get("finance_block", "defaulted")),
        ("wdv_rate", "Tax depreciation (WDV) rate", a.tax_wdv_rate, "fraction", a.sources.get("finance_block", "defaulted")),
        ("wdv_eligible", "WDV-eligible fraction of depreciable cost", a.tax_wdv_eligible_fraction, "fraction", a.sources.get("finance_block", "defaulted")),
        ("book_years", "Book depreciation years", a.book_depreciation_years, "years", a.sources.get("finance_block", "defaulted")),
        ("residual", "Residual value fraction", a.residual_value_fraction, "fraction", a.sources.get("finance_block", "defaulted")),
        ("discount", "Discount rate", a.discount_rate, "per year", a.sources.get("finance_block", "defaulted")),
        ("target_irr", "Target equity IRR", a.target_equity_irr, "per year", a.sources.get("finance_block", "defaulted")),
        ("min_dscr", "Minimum DSCR", a.min_dscr, "x", a.sources.get("finance_block", "defaulted")),
        ("insurance", "Insurance (fraction of hard cost per year)", a.insurance_fraction_of_capex, "fraction", a.sources.get("finance_block", "defaulted")),
        ("tariff_esc", "Tariff escalation", a.tariff_escalation, "per year", "tender rules"),
        ("terminal", "Terminal value", a.terminal_value_inr, "INR", a.sources.get("finance_block", "defaulted")),
    ]
    for i, (key, label, value, unit, source) in enumerate(scalar, start=2):
        inp.cell(i, 1, label)
        c = inp.cell(i, 2, value)
        c.font = INPUT_FONT
        c.fill = INPUT_FILL
        inp.cell(i, 3, unit)
        inp.cell(i, 4, source)
        _name(wb, key, f"Inputs!$B${i}")
    row = len(scalar) + 3
    inp.cell(row, 1, "Capex items").font = HEADER_FONT
    inp.cell(row, 2, "INR")
    inp.cell(row, 3, "Tax block")
    inp.cell(row, 4, "Source")
    capex_start = row + 1
    for j, item in enumerate(a.capex):
        r = capex_start + j
        inp.cell(r, 1, item.name)
        c = inp.cell(r, 2, item.amount_inr)
        c.font = INPUT_FONT
        c.fill = INPUT_FILL
        inp.cell(r, 3, item.tax_block)
        inp.cell(r, 4, item.source)
    capex_end = capex_start + max(len(a.capex), 1) - 1
    _name(wb, "capex_items", f"Inputs!$B${capex_start}:$B${capex_end}")
    _name(wb, "capex_blocks", f"Inputs!$C${capex_start}:$C${capex_end}")
    row = capex_end + 2
    inp.cell(row, 1, "Fixed opex items").font = HEADER_FONT
    inp.cell(row, 2, "INR per year (year 1)")
    inp.cell(row, 3, "Escalation per year")
    inp.cell(row, 4, "Source")
    opex_start = row + 1
    for j, item in enumerate(a.opex):
        r = opex_start + j
        inp.cell(r, 1, item.name)
        c = inp.cell(r, 2, item.amount_inr_per_year)
        c.font = INPUT_FONT
        c.fill = INPUT_FILL
        e = inp.cell(r, 3, item.escalation)
        e.font = INPUT_FONT
        e.fill = INPUT_FILL
        inp.cell(r, 4, item.source)
    opex_end = opex_start + max(len(a.opex), 1) - 1
    row = opex_end + 2
    inp.cell(row, 1, "Augmentation / replacement (INR by year)").font = HEADER_FONT
    aug_row = row + 1
    inp.cell(aug_row, 1, "Year")
    for y in range(1, n + 1):
        inp.cell(aug_row, 1 + y, y)
    inp.cell(aug_row + 1, 1, "Amount INR")
    aug_by_year = {y: 0.0 for y in range(1, n + 1)}
    for y, amt in a.augmentation:
        if 1 <= y <= n:
            aug_by_year[y] += amt
    for y in range(1, n + 1):
        c = inp.cell(aug_row + 1, 1 + y, aug_by_year[y])
        c.font = INPUT_FONT
        c.fill = INPUT_FILL
    inp.column_dimensions["A"].width = 44
    inp.column_dimensions["D"].width = 50
    # ---------------- Dispatch (fixed snapshot)
    dsp = wb.create_sheet("Dispatch")
    dsp.append(["FIXED RUN SNAPSHOT: these values come from the optimisation run and do not recalculate. A change to capacities, profiles, losses or rules needs a new application run.", ""])
    dsp["A1"].font = Font(bold=True, color="8B0000")
    headers = ["Year", "PPA MWh", "Merchant MWh", "Merchant revenue INR", "Purchase MWh", "Purchase cost INR", "Penalty MWh", "Fuel cost INR", "Thermal variable cost INR", "Start-up cost INR", "External green MWh"]
    dsp.append(headers)
    for c in dsp[2]:
        c.font = HEADER_FONT
    for i, o in enumerate(operating[:n], start=1):
        dsp.append([i, o.ppa_mwh, o.merchant_mwh, o.merchant_revenue_inr, o.purchase_mwh, o.purchase_cost_inr, o.penalty_mwh, o.fuel_cost_inr, o.thermal_variable_cost_inr, o.startup_cost_inr, o.external_green_mwh])
        for c in dsp[2 + i]:
            c.fill = SNAPSHOT_FILL
    # ---------------- Capex & funding
    cap = wb.create_sheet("Capex & Funding")
    cap.append(["Item", "Formula", "Value INR"])
    cap_rows = [
        ("Hard cost", "=SUM(capex_items)"),
        ("Contingency", "=C2*contingency"),
        ("Base cost", "=C2+C3"),
        ("IDC", "=idc"),
        ("Fees", "=fees"),
        ("Total project cost", "=C4+C5+C6"),
        ("Debt", "=C7*debt_fraction"),
        ("Equity", "=C7-C8"),
        ("Land (non-depreciable)", '=SUMIF(capex_blocks,"land_non_depreciable",capex_items)'),
        ("Depreciable cost", "=C7-C10"),
        ("DSRA at COD (months of year-1 debt service)", "='Debt Tax Cashflow'!B16*dsra_months/12"),
    ]
    for i, (label, f) in enumerate(cap_rows, start=2):
        cap.cell(i, 1, label)
        cap.cell(i, 2, f)
        cap.cell(i, 3, f)
    for key, r in (("hard_cost", 2), ("total_cost", 7), ("debt", 8), ("equity", 9), ("land", 10), ("depreciable", 11), ("dsra", 12)):
        _name(wb, key, f"'Capex & Funding'!$C${r}")
    # ---------------- Operations (annual statement)
    ops = wb.create_sheet("Operations")
    labels = ["Year", "Tariff escalator", "PPA revenue", "Merchant revenue", "Penalties", "Purchases", "Fuel and variable", "Fixed opex", "Insurance", "Receivables", "WC interest", "Revenue net of penalties", "EBITDA", "Delta working capital", "Augmentation"]
    for i, lab in enumerate(labels, start=1):
        ops.cell(i, 1, lab).font = HEADER_FONT
    for y in range(1, n + 1):
        col = get_column_letter(1 + y)
        d = 2 + y  # dispatch row
        ops[f"{col}1"] = y
        ops[f"{col}2"] = f"=(1+tariff_esc)^({col}1-1)"
        ops[f"{col}3"] = f"=Dispatch!B{d}*MAX(0,tariff-psm)*1000*{col}2"
        ops[f"{col}4"] = f"=Dispatch!D{d}"
        ops[f"{col}5"] = f"=Dispatch!G{d}*penalty_mult*tariff*1000*{col}2"
        ops[f"{col}6"] = f"=Dispatch!F{d}"
        ops[f"{col}7"] = f"=Dispatch!H{d}+MAX(0,Dispatch!I{d}-Dispatch!H{d})+Dispatch!J{d}"
        ops[f"{col}8"] = "=" + ("+".join(f"Inputs!$B${opex_start + j}*(1+Inputs!$C${opex_start + j})^({col}$1-1)" for j in range(len(a.opex))) or "0")
        ops[f"{col}9"] = "=hard_cost*insurance"
        ops[f"{col}10"] = f"=MAX({col}12,0)*receivable_days/365"
        ops[f"{col}11"] = f"={col}10*wc_rate"
        ops[f"{col}12"] = f"={col}3+{col}4-{col}5"
        ops[f"{col}13"] = f"={col}12-{col}6-{col}7-{col}8-{col}9-{col}11"
        prev = get_column_letter(y) if y > 1 else None
        ops[f"{col}14"] = f"={col}10-{prev}10" if prev else f"={col}10"
        ops[f"{col}15"] = f"=Inputs!{get_column_letter(1 + y)}{aug_row + 1}"
    ops.column_dimensions["A"].width = 26
    # ---------------- Debt, depreciation, tax
    fin = wb.create_sheet("Debt Tax Cashflow")
    labels2 = [
        "Year", "Debt opening", "Interest", "CFADS (pre-tax proxy for sculpting)", "Principal", "Debt closing", "Book depreciation", "WDV opening", "Tax depreciation", "Taxable income before loss", "Loss brought forward",
        "Taxable after loss", "Tax", "PAT", "CFADS", "Debt service", "DSCR", "Terminal and DSRA release", "FCFE", "Project cash flow",
    ]
    for i, lab in enumerate(labels2, start=1):
        fin.cell(i, 1, lab).font = HEADER_FONT
    for y in range(1, n + 1):
        col = get_column_letter(1 + y)
        prev = get_column_letter(y) if y > 1 else None
        fin[f"{col}1"] = y
        fin[f"{col}2"] = "=debt" if y == 1 else f"={prev}6"
        fin[f"{col}3"] = f"={col}2*interest"
        # tax for sculpting uses EBITDA - interest - tax depreciation (same-year), consistent with the converged backend
        fin[f"{col}4"] = f"=Operations!{col}13-Operations!{col}14-{col}13"
        fin[f"{col}5"] = f"=IF({col}1>tenor,0,IF({col}1=tenor,{col}2,MIN({col}2,MAX(0,{col}4/target_dscr-{col}3))))"
        fin[f"{col}6"] = f"={col}2-{col}5"
        fin[f"{col}7"] = f"=depreciable*(1-residual)/book_years" + (f"+SUM(Operations!$B$15:{prev}15)*(1-residual)/(ppa_years-{col}1+1)" if prev else "")
        fin[f"{col}8"] = "=depreciable*wdv_eligible" if y == 1 else f"={prev}8*(1-wdv_rate)+Operations!{col}15*wdv_eligible"
        fin[f"{col}9"] = f"={col}8*wdv_rate+depreciable*(1-wdv_eligible)/book_years"
        fin[f"{col}10"] = f"=Operations!{col}13-{col}3-{col}9"
        fin[f"{col}11"] = "=0" if y == 1 else f"=MAX(0,-{prev}12)"
        fin[f"{col}12"] = f"={col}10-{col}11"
        fin[f"{col}13"] = f"=MAX(0,{col}12)*tax_rate"
        fin[f"{col}14"] = f"=Operations!{col}13-{col}7-{col}3-{col}13"
        fin[f"{col}15"] = f"=Operations!{col}13-{col}13-Operations!{col}14"
        fin[f"{col}16"] = f"={col}3+{col}5"
        fin[f"{col}17"] = f'=IF({col}16>0,{col}15/{col}16,"n/a")'
        fin[f"{col}18"] = "=terminal+dsra" if y == n else "=0"
        fin[f"{col}19"] = f"={col}15-{col}16-Operations!{col}15+{col}18"
        fin[f"{col}20"] = f"={col}15-Operations!{col}15" + ("+terminal" if y == n else "")
    fin.column_dimensions["A"].width = 36
    last = get_column_letter(1 + n)
    # ---------------- Cash flows and returns
    ret = wb.create_sheet("Returns")
    cy = len(result.construction_equity_flows)
    ret.append(["Period", *[f"C{-(cy - i)}" for i in range(cy)], *[f"Y{y}" for y in range(1, n + 1)]])
    eq = ["Equity cash flow"]
    pj = ["Project cash flow"]
    for i in range(cy):
        share = a.drawdown_profile[i] if i < len(a.drawdown_profile) else 1.0 / cy
        eq.append(f"=-equity*{share}" + ("-dsra" if i == cy - 1 else ""))
        pj.append(f"=-(total_cost-fees)*{share}")
    for y in range(1, n + 1):
        col = get_column_letter(1 + y)
        eq.append(f"='Debt Tax Cashflow'!{col}19")
        pj.append(f"='Debt Tax Cashflow'!{col}20")
    ret.append(eq)
    ret.append(pj)
    end_col = get_column_letter(1 + cy + n)
    ret.append([])
    ret.append(["Equity IRR", f"=IRR(B2:{end_col}2)"])
    ret.append(["Project IRR", f"=IRR(B3:{end_col}3)"])
    ret.append(["NPV equity at discount rate", f"=NPV(discount,C2:{end_col}2)+B2"])
    ret.append(["NPV project at discount rate", f"=NPV(discount,C3:{end_col}3)+B3"])
    ret.append(["Minimum DSCR", f"=MIN('Debt Tax Cashflow'!B17:{last}17)"])
    ret.append(["Average DSCR", f"=AVERAGE('Debt Tax Cashflow'!B17:{last}17)"])
    ret.append(["Meets target (IRR and DSCR)", "=AND(B5>=target_irr,OR(debt<=0,B9>=min_dscr))"])
    ret.append(["Total project cost", "=total_cost"])
    for key, r in (("equity_irr", 5), ("project_irr", 6), ("npv_equity", 7), ("npv_project", 8), ("min_dscr_calc", 9)):
        _name(wb, key, f"Returns!$B${r}")
    ret.column_dimensions["A"].width = 30
    # ---------------- Sensitivities (formula-driven: tariff and capex deltas via data table style re-computation is not possible without re-running the sheets, so the sheet shows elasticities from the backend plus live tariff sensitivity of year-1 revenue)
    sen = wb.create_sheet("Sensitivities")
    sen.append(["Case", "Equity IRR (backend, at export)", "Project IRR", "NPV equity (cr)", "Min DSCR", "Note"])
    for s in provenance.get("sensitivities", []):
        sen.append([s["case"], s.get("equity_irr"), s.get("project_irr"), s.get("npv_equity_cr"), s.get("min_dscr"), "backend finance.v1 sensitivity on the fixed dispatch"])
    sen.append([])
    sen.append(["Live: year-1 PPA revenue at tariff +5%", "=Dispatch!B3*MAX(0,tariff*1.05-psm)*1000", "", "", "", "formula"])
    sen.append(["Live: year-1 PPA revenue at tariff -5%", "=Dispatch!B3*MAX(0,tariff*0.95-psm)*1000", "", "", "", "formula"])
    sen.column_dimensions["A"].width = 44
    # ---------------- Checks
    chk = wb.create_sheet("Checks")
    chk.append(["Check", "Workbook", "Backend at export", "Difference", "Status"])
    checks = [
        ("Total project cost INR", "=total_cost", result.total_project_cost_inr),
        ("Debt INR", "=debt", result.debt_inr),
        ("Year-1 PPA revenue INR", "=Operations!B3", float(result.statement.revenue_ppa[0])),
        ("Year-1 EBITDA INR", "=Operations!B13", float(result.statement.ebitda[0])),
        ("Year-1 CFADS INR", "='Debt Tax Cashflow'!B15", float(result.statement.cfads[0])),
        ("Year-1 interest INR", "='Debt Tax Cashflow'!B3", float(result.statement.interest[0])),
        ("Equity IRR", "=equity_irr", result.equity_irr if result.equity_irr is not None else 0.0),
        ("Project IRR", "=project_irr", result.project_irr if result.project_irr is not None else 0.0),
        ("NPV equity INR", "=npv_equity", result.npv_equity_inr),
        ("Minimum DSCR", "=min_dscr_calc", result.min_dscr if result.min_dscr != float("inf") else 0.0),
    ]
    for i, (label, f, backend) in enumerate(checks, start=2):
        chk.cell(i, 1, label)
        chk.cell(i, 2, f)
        c = chk.cell(i, 3, backend)
        c.fill = SNAPSHOT_FILL
        chk.cell(i, 4, f"=B{i}-C{i}")
        chk.cell(i, 5, f'=IF(ABS(D{i})<=MAX(ABS(C{i})*0.001,IF(ABS(C{i})<10,0.0001,1)),"OK","DIFFERS (inputs changed, or see Read Me simplifications)")')
        chk.cell(i, 5).fill = CHECK_FILL
    chk.append([])
    chk.append(["Audit references", f"run {provenance.get('run_id')} · inputs {provenance.get('inputs_hash')} · rules {provenance.get('rules_version')} · exported {datetime.now(UTC).isoformat()}"])
    chk.column_dimensions["A"].width = 30
    chk.column_dimensions["E"].width = 60
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("=") and cell.fill != INPUT_FILL:
                    cell.font = FORMULA_FONT
                cell.alignment = Alignment(vertical="top", wrap_text=False)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def recalculate_with_libreoffice(xlsx_bytes: bytes, timeout: int = 120) -> dict[str, Any]:
    """Recalculate with LibreOffice headless (the only spreadsheet engine available here) and
    read the computed values back. Returns {ok, values: {sheet: {cell: value}}, error}."""
    import shutil
    import subprocess
    import tempfile
    from pathlib import Path

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return {"ok": False, "error": "LibreOffice (soffice) not found: recalculation verification blocked", "values": {}}
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "model.xlsx"
        src.write_bytes(xlsx_bytes)
        out = Path(tmp) / "out"
        out.mkdir()
        try:
            subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(src)], check=True, capture_output=True, timeout=timeout, env={"HOME": tmp, "PATH": "/usr/bin:/bin"})
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
            return {"ok": False, "error": f"LibreOffice recalculation failed: {exc}", "values": {}}
        from openpyxl import load_workbook

        wb = load_workbook(out / "model.xlsx", data_only=True)
        values = {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None} for ws in wb.worksheets}
        return {"ok": True, "values": values, "error": None}


def verify(xlsx_bytes: bytes, result: FinanceResult, *, rel: float = 1e-3) -> dict[str, Any]:
    """Recalculate and compare the workbook's results with finance.v1."""
    recalc = recalculate_with_libreoffice(xlsx_bytes)
    if not recalc["ok"]:
        return {"status": "blocked", "error": recalc["error"], "comparisons": []}
    ret = recalc["values"].get("Returns", {})
    chk = recalc["values"].get("Checks", {})
    comparisons = []
    for label, cell, backend in (("equity_irr", "B5", result.equity_irr), ("project_irr", "B6", result.project_irr), ("npv_equity", "B7", result.npv_equity_inr), ("min_dscr", "B9", result.min_dscr)):
        wbv = ret.get(cell)
        ok = wbv is not None and backend is not None and isinstance(wbv, (int, float)) and abs(float(wbv) - float(backend)) <= max(rel * abs(float(backend)), 1e-4 if abs(float(backend)) < 10 else 1.0)
        comparisons.append({"metric": label, "workbook": wbv, "backend": backend, "ok": bool(ok)})
    statuses = [v for k, v in chk.items() if k.startswith("E") and isinstance(v, str) and k != "E1"]
    return {"status": "verified" if all(c["ok"] for c in comparisons) else "mismatch", "comparisons": comparisons, "checks_sheet": statuses, "error": None}
